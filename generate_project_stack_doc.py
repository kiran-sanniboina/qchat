from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PROJECT_STACKS_AND_MODULES.md"
OUTPUT = ROOT / "PROJECT_STACKS_AND_MODULES.docx"


def set_cell_shading(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shading = props.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        props.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for index, value in enumerate(rows[0]):
        set_cell_text(table.rows[0].cells[index], value, bold=True, color=(255, 255, 255))
        set_cell_shading(table.rows[0].cells[index], "1F4E79")
    for row in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if len(table.rows) % 2 == 0:
                set_cell_shading(cells[index], "F3F6F9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_run_with_code(paragraph, text):
    parts = text.split("`")
    for index, part in enumerate(parts):
        run = paragraph.add_run(part)
        run.font.name = "Aptos Mono" if index % 2 else "Aptos"
        run._element.rPr.rFonts.set(qn("w:ascii"), run.font.name)
        run._element.rPr.rFonts.set(qn("w:hAnsi"), run.font.name)
        run.font.size = Pt(10)
        if index % 2:
            run.font.color.rgb = RGBColor(31, 78, 121)


def render_markdown(doc, lines):
    index = 0
    while index < len(lines):
        line = lines[index].rstrip("\n")
        if not line.strip():
            index += 1
            continue

        if line.startswith("| "):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                raw = lines[index].strip()
                cells = [cell.strip() for cell in raw.strip("|").split("|")]
                if all(set(cell) <= {"-", ":", " "} for cell in cells):
                    index += 1
                    continue
                rows.append(cells)
                index += 1
            add_table(doc, rows)
            continue

        if line.startswith("# "):
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(line[2:].strip())
            run.font.name = "Aptos Display"
            run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
            run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
            run.font.color.rgb = RGBColor(31, 78, 121)
            index += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue

        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_run_with_code(paragraph, line[2:].strip())
            index += 1
            continue

        if len(line) > 2 and line[0].isdigit() and line[1:3] == ". ":
            paragraph = doc.add_paragraph(style="List Number")
            add_run_with_code(paragraph, line[3:].strip())
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        add_run_with_code(paragraph, line.strip())
        index += 1


def set_document_defaults(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10)

    for style_name, size, color in (("Heading 1", 16, (31, 78, 121)), ("Heading 2", 12, (47, 84, 150))):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("QChat - Stacks and Modules")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)


def main():
    doc = Document()
    set_document_defaults(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    render_markdown(doc, lines)
    doc.core_properties.title = "QChat: Stacks and Modules"
    doc.core_properties.subject = "Technology stack, project modules, and reasons for use"
    doc.core_properties.author = "Codex"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
